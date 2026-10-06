from pathlib import Path
from bs4 import BeautifulSoup
from playwright.sync_api import sync_playwright
from pptx import Presentation
import pandas as pd,numpy as np,json,zipfile,hashlib,re
P=Path(__file__).resolve().parents[1];O=P/'output';D=P/'data';report=O/'国证价值100_深度研究.html'
q={};doc=report.read_text();s=BeautifulSoup(doc,'html.parser');ids=[n['id'] for n in s.select('[id]')]
q['html_sections']=len(s.select('section'));q['figures']=len(s.select('figure'));q['unique_ids']=len(ids)==len(set(ids));q['missing_internal_targets']=[a['href'] for a in s.select('a[href^="#"]') if a['href'][1:] not in ids]
q['missing_local_downloads']=[a['href'] for a in s.select('a[download]') if not (O/a['href']).exists()]
q['key_assertions']=all(x in s.get_text() for x in ['8.3532','47.65','23.93%','13.34%','6.33%','3539','2024-10-29','尚未'])
q['embedded_csv_matches']=all(json.loads(s.select_one('#embedded-data').string)[f]==(D/f).read_text(encoding='utf-8-sig') for f in json.loads(s.select_one('#embedded-data').string))
tr=pd.read_csv(D/'domestic_total_return_daily.csv',index_col=0,parse_dates=True);perf=pd.read_csv(D/'performance_metrics.csv');row=perf[(perf.name=='价值100')&(perf.period=='full_backtest_mixed')].iloc[0];ret=tr['价值100'].iloc[-1]/tr['价值100'].iloc[0]-1
assert abs(ret-row.total)<1e-12
q['numeric_independent_checks']={'total_return_matches':True,'common_rows':len(tr),'nulls':int(tr.isna().sum().sum()),'duplicate_dates':int(tr.index.duplicated().sum()),'constituents':len(pd.read_csv(D/'constituents_20260930.csv'))}
with sync_playwright() as pw:
 browser=pw.chromium.launch(executable_path='/usr/bin/chromium',headless=True,args=['--no-sandbox','--disable-dev-shm-usage'])
 for width,height,label in [(390,844,'手机'),(1440,950,'桌面')]:
  page=browser.new_page(viewport={'width':width,'height':height},accept_downloads=True);errors=[];page.on('pageerror',lambda er:errors.append(str(er)))
  page.set_content(doc,wait_until='load')
  page.screenshot(path=str(O/('QA_'+label+'首屏.png')),full_page=False)
  q[label+'_overflow']=page.evaluate('document.documentElement.scrollWidth>innerWidth+1')
  page.locator('a[href="#verdict"]').first.click();page.screenshot(path=str(O/('QA_'+label+'结论.png')))
  page.select_option('#group-filter','A');assert page.locator('#competitor-table tbody tr:visible').count()==5
  page.fill('#competitor-search','现金流');assert 0<page.locator('#competitor-table tbody tr:visible').count()<=5
  page.select_option('#group-filter','all');page.fill('#competitor-search','');assert page.locator('#competitor-table tbody tr:visible').count()==12
  page.locator('[data-roll="10"]').click();assert page.locator('#roll-10').is_visible() and not page.locator('#roll-1').is_visible()
  page.locator('#chart-recent .zoom').click();assert page.locator('#chart-dialog').is_visible();page.locator('.close-dialog').click()
  with page.expect_download() as di:page.locator('[data-file="scenario_returns.csv"]').click()
  download=di.value;assert download.suggested_filename=='scenario_returns.csv';assert Path(download.path()).read_text(encoding='utf-8-sig')==(D/'scenario_returns.csv').read_text(encoding='utf-8-sig')
  page.locator('#competition').scroll_into_view_if_needed();page.screenshot(path=str(O/('QA_'+label+'竞品.png')))
  page.locator('#current').scroll_into_view_if_needed();page.screenshot(path=str(O/('QA_'+label+'估值.png')))
  q[label+'_page_errors']=errors;q[label+'_controls']='filters, rolling tabs, zoom, CSV download passed';page.close()
 browser.close()
pr=Presentation(O/'国证价值100_结论版.pptx');q['pptx_slides']=len(pr.slides);overflow=[]
for n,sl in enumerate(pr.slides,1):
 for sh in sl.shapes:
  if sh.left<0 or sh.top<0 or sh.left+sh.width>pr.slide_width+1000 or sh.top+sh.height>pr.slide_height+1000:overflow.append((n,sh.name))
q['pptx_geometry_overflow']=overflow
for f in ['国证价值100_数据与计算.zip','国证价值100_结论版.pptx','国证价值100_核对工作簿.xlsx']:
 with zipfile.ZipFile(O/f) as z:assert z.testzip() is None
q['zip_office_integrity']='passed';q['workbook_sheets']=len(pd.ExcelFile(O/'国证价值100_核对工作簿.xlsx').sheet_names)
q['qa_note']='PPTX checked as OOXML and geometry; no LibreOffice installed for native full rendering. HTML actually rendered with Chromium set_content on mobile and desktop; environment policy blocks file:// navigation.'
(D/'qa_results.json').write_text(json.dumps(q,ensure_ascii=False,indent=2));print(json.dumps(q,ensure_ascii=False,indent=2))
assert q['unique_ids'] and not q['missing_internal_targets'] and not q['missing_local_downloads'] and q['key_assertions'] and q['embedded_csv_matches'] and not q['手机_overflow'] and not q['桌面_overflow'] and not q['手机_page_errors'] and not q['桌面_page_errors'] and not overflow
